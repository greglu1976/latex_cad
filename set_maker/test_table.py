# Создание шаблона разделов на основе шаблона origin.docx

from docxtpl import DocxTemplate
from tables import add_table_settings
from docx import Document

from add_sect_binaries import add_sect_binaries, add_sect_reg, add_sect_mtrx_outs
from docx_handler import add_new_section, add_new_section_landscape, add_new_section_test

doc = Document('origin.docx')

# Добавляем подразделы с дискр вх/вых в существующий раздел Конфигурация
add_sect_binaries(doc, 'Модули дискретных входов', 'bin_input')
add_sect_binaries(doc, 'Модули дискретных выходов', 'bin_output')

#============================== УСТАВКИ РЗА ==================================
add_new_section(doc) # Создаем раздел для РЗА
section_relay = doc.add_paragraph(f'Уставки РЗА')
section_relay.style = 'ДОК Заголовок 1'
# Вставляем шаблон таблицы для уставок РЗА
add_sect_binaries(doc, 'Группа уставок №1', 'rza_func')

#========================= РЕГИСТРАЦИЯ =======================================
add_new_section_landscape(doc) # Создаем раздел для регистрации
# Добавляем заголовок
section_reg = doc.add_paragraph('Настройка регистрации')
section_reg.style = 'ДОК Заголовок 1'
add_sect_reg(doc)

#=================================МАТРИЦА ВЫХОДОВ ================================
add_new_section_landscape(doc) # Создаем раздел для матрицы вх/вых
# Добавляем заголовок
section_mtrx = doc.add_paragraph('Матрица дискретных входов и выходных реле')
section_mtrx.style = 'ДОК Заголовок 1'
add_sect_mtrx_outs(doc)

doc.save('templ2.docx')