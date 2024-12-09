from docx.shared import Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import parse_xml

table_settings = (
    Inches(0.5),  # №
    Inches(2.5),  # Описание
    Inches(2),    # Наименование
    Inches(0.5),  # -
    Inches(3),    # Значение/Диапазон
    Inches(0.5),  # Ед. изм.
    Inches(0.5),  # Шаг
    Inches(1),    # Значение по умолчанию
    Inches(1)     # Уставка
)

def set_cell_vertical_alignment(cell, align="center"):
    """Установка вертикального выравнивания в ячейке"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcVAlign = parse_xml(f'<w:vAlign xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="{align}"/>')
    tcPr.append(tcVAlign)

def set_cell_border(cell, **kwargs):
    """Установка границ ячейки"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    
    for key, value in kwargs.items():
        tag = 'w:{}'.format(key)
        element = parse_xml('<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                          '<{} w:val="{}" w:sz="4" w:space="0" w:color="auto"/>'
                          '</w:tcBorders>'.format(tag, value.get("val", "single")))
        tcPr.append(element)

def set_repeat_table_header(row):
    """Установка повторения заголовка таблицы на следующих страницах"""
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = parse_xml('<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="true"/>')
    trPr.append(tblHeader)

def add_table_settings(doc):
    # Создаем таблицу
    table = doc.add_table(rows=2, cols=9)
    
    # Устанавливаем стиль и отключаем автоподбор
    table.style = 'Стиль2'
    table.allow_autofit = False
    
    # Устанавливаем фиксированный макет таблицы с правильным пространством имен
    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )
    
    # Устанавливаем ширину столбцов
    for column_index, width in enumerate(table_settings):
        table.columns[column_index].width = width

    # Заполняем первую строку заголовка
    hdr_cells = table.rows[0].cells
    headers = ['№', 'Описание', 'Наименование', '', 'Значение / Диапазон', 
              'Ед. изм.', 'Шаг', 'Значение по умолчанию', 'Уставка']
    
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Устанавливаем повторение заголовка
    set_repeat_table_header(table.rows[0])

    # Заполняем вторую строку заголовка
    hdr_cells = table.rows[1].cells
    hdr_cells[2].text = 'ПО'
    hdr_cells[3].text = 'ФСУ'

    # Устанавливаем повторение второй строки заголовка
    set_repeat_table_header(table.rows[1])

    # Форматируем вторую строку заголовка
    for i in range(9):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_border(hdr_cells[i], bottom={"val": "double"})

    # Объединяем ячейки в первой строке
    table.cell(0, 2).merge(table.cell(0, 3))

    # Устанавливаем общую ширину таблицы
    table.width = sum(table_settings)

    return table
